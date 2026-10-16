"""Genera la matriz UX en Excel. Las columnas de emoción, valoración y evidencia salen del CSV real.
Uso: python scripts/matriz.py"""
import sys
from collections import Counter
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from util import cargar, matriz_valores
from config_momentos import MOMENTOS

ra = sys.argv[1] if len(sys.argv) > 1 else "datos/pruebas_antes.csv"
rd = sys.argv[2] if len(sys.argv) > 2 else "datos/pruebas_despues.csv"
salida = sys.argv[3] if len(sys.argv) > 3 else "matriz/matriz_ux.xlsx"

a = cargar(ra)
try:
    d = cargar(rd)
except SystemExit:
    d = None
    print("Aviso. Todavía no hay datos válidos del después. La columna de validación queda pendiente.")

def textos(df, mid, col):
    s = df[(df["momento"] == mid)][col].dropna().astype(str).str.strip()
    s = [t for t in s if t]
    return s

wb = Workbook(); ws = wb.active; ws.title = "Matriz UX"
enc = ["Momento", "Fase", "Emoción reportada", "Valoración promedio", "Valoraciones individuales", "ID de captura",
       "Evidencia observable", "Mejora propuesta", "Principio de IHC", "Resultado de la validación"]
ws.append(enc)
va = matriz_valores(a); vd = matriz_valores(d) if d is not None else None
for mo in MOMENTOS:
    i = mo["id"]
    emos = Counter(t.lower() for t in textos(a, i, "emocion_reportada") if "emocion_reportada" in a.columns)
    emo_txt = ", ".join(f"{k} {v}" if v > 1 else k for k, v in emos.most_common()) or "Sin dato"
    indiv = ", ".join(f"{p} {int(v):+d}" for p, v in va[i].items())
    evid = textos(a, i, "evidencia_observada") + [f"Dijo {t}" for t in textos(a, i, "comentario")]
    evid_txt = ". ".join(dict.fromkeys(evid)) or "Sin dato"
    if vd is not None:
        pa_, pd_ = va[i].mean(), vd[i].mean()
        val = f"Antes {pa_:+.1f}, después {pd_:+.1f}, cambio {pd_ - pa_:+.1f}"
    else:
        val = "Pendiente de la prueba del rediseño"
    ws.append([f"M{i} {mo['nombre']}", mo["fase"], emo_txt, round(va[i].mean(), 2), indiv, mo["captura"],
               evid_txt, mo["mejora"], mo["principio"], val])
anchos = [30, 11, 26, 14, 28, 11, 48, 44, 44, 32]
fino = Side(style="thin", color="B7C0D0")
for c, w in zip("ABCDEFGHIJ", anchos): ws.column_dimensions[c].width = w
for r in ws.iter_rows():
    for c in r:
        c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(top=fino, bottom=fino, left=fino, right=fino)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
ws.freeze_panes = "B2"
for nombre, df in (("Datos antes", a), ("Datos después", d)):
    if df is None: continue
    w2 = wb.create_sheet(nombre); w2.append(list(df.columns))
    for row in df.itertuples(index=False): w2.append(list(row))
    for c in w2[1]: c.font = Font(bold=True)
wb.save(salida); print("Guardado", salida)
