"""Arma el informe final en Word a partir de los CSV reales, la matriz y las imágenes.
Uso desde la raíz del repositorio. python scripts/informe.py"""
import os, sys
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from openpyxl import load_workbook
from util import cargar, matriz_valores, momentos_criticos, validar_mismos_participantes, NOMBRES
from config_momentos import MOMENTOS, TAREA

NAVY = "1F3864"
SALIDA = "informe/informe_final.docx"
a = cargar("datos/pruebas_antes.csv"); d = cargar("datos/pruebas_despues.csv")
validar_mismos_participantes(a, d)
va, vd = matriz_valores(a), matriz_valores(d); pa, pdv = va.mean(), vd.mean()
bajo, caida, delta = momentos_criticos(pa)
npart = va.shape[0]

equipo = pd.read_csv("datos/equipo.csv").fillna("")
enl = {}
for linea in open("datos/enlaces.txt", encoding="utf-8"):
    if "=" in linea:
        k, v = linea.strip().split("=", 1); enl[k] = v

def textos(df, mid, col):
    if col not in df.columns: return []
    s = df[df["momento"] == mid][col].dropna().astype(str).str.strip()
    return [t for t in s if t]

doc = Document(); s = doc.sections[0]; s.page_width = Cm(21); s.page_height = Cm(29.7)
s.left_margin = s.right_margin = Cm(2); s.top_margin = Cm(2.2); s.bottom_margin = Cm(1.8)
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
for n, sz in (("Heading 1", 15), ("Heading 2", 12)):
    h = doc.styles[n]; h.font.name = "Arial"; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = RGBColor.from_string(NAVY)
    for k in ("w:eastAsia", "w:ascii", "w:hAnsi"): h.element.rPr.rFonts.set(qn(k), "Arial")
    h.paragraph_format.space_before = Pt(14); h.paragraph_format.space_after = Pt(6)
fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
def fld(p, code):
    r = p.add_run(); r.font.size = Pt(8)
    for t, txt in (("begin", None), (None, code), ("end", None)):
        if t: e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
        else: e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
        r._r.append(e)
r0 = fp.add_run("Grupo C, Interacción Humano Computador. Página "); r0.font.size = Pt(8); fld(fp, "PAGE")

def P(text, bold=False, size=None, after=6, align=None, color=None, italic=False):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.15
    if align is not None: p.alignment = align
    return p
def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color); tcPr.append(sh)
def table(heads, rows, widths, size=8.5, first_bold=True):
    t = doc.add_table(rows=1, cols=len(heads)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for c, h, w in zip(t.rows[0].cells, heads, widths):
        c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255); shade(c, NAVY); c.width = Cm(w)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, (c, v, w) in enumerate(zip(cells, row, widths)):
            c.text = ""; r = c.paragraphs[0].add_run(str(v)); r.font.size = Pt(size); r.bold = (i == 0 and first_bold); c.width = Cm(w)
            if ri % 2: shade(c, "F3F6FB")
    for i, w in enumerate(widths): t.columns[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2); return t
def IMG(path, width, cap):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(path, width=Cm(width))
    p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(2)
    P(cap, italic=True, size=9, after=10, align=WD_ALIGN_PARAGRAPH.CENTER, color="555555")

# portada
P("Universidad Técnica de Ambato", True, 18, 2, WD_ALIGN_PARAGRAPH.CENTER, NAVY)
P("Facultad de Ingeniería en Sistemas, Electrónica e Industrial", False, 11, 2, WD_ALIGN_PARAGRAPH.CENTER)
P("Carrera de Software", False, 11, 60, WD_ALIGN_PARAGRAPH.CENTER)
P("Evaluación y rediseño de la experiencia emocional del usuario", True, 24, 8, WD_ALIGN_PARAGRAPH.CENTER, NAVY)
P("Pantalla B-4 Mapa del flujo de compra de SuperOfertas Ecuador", False, 13, 50, WD_ALIGN_PARAGRAPH.CENTER, "444444")
integ = " ".join(f"{r.nombre}, rol {r.rol}." for r in equipo.itertuples())
cov = [("Asignatura", "Interacción Humano Computador"), ("Docente", "Ing. José Caiza"), ("Grupo", "Grupo C"),
       ("Integrantes", integ), ("Fecha", enl.get("fecha", "")), ("Figma", enl.get("figma", "")), ("Repositorio", enl.get("repositorio", ""))]
t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for k, v in cov:
    c = t.add_row().cells; c[0].text = ""; c[1].text = ""
    r = c[0].paragraphs[0].add_run(k); r.bold = True; r.font.size = Pt(11); shade(c[0], "EEF2FA")
    r = c[1].paragraphs[0].add_run(v); r.font.size = Pt(11); c[0].width = Cm(3.5); c[1].width = Cm(13.5)
doc.add_page_break()

doc.add_heading("1. Caso de estudio y tarea", 1)
P("Analizamos la pantalla del mapa, paso 4 de 5, del flujo de compra de SuperOfertas Ecuador. Es un prototipo ficticio del Grupo 1 de la clase. La pantalla pide revisar la dirección de entrega y tiene un banner rojo de oferta, una cuenta regresiva, un botón Guardar dirección en gris y un botón OMITIR muy grande.")
P("La tarea que hicieron los participantes fue la misma antes y después. " + TAREA)
doc.add_heading("2. Método y participantes", 1)
P(f"Probamos con {npart} personas ajenas al equipo, identificadas como {', '.join(va.index)}. Se les pidió permiso verbal y no se guardan nombres ni fotos de rostros. En cada uno de los cinco momentos dijeron un número de -3 muy negativa a +3 muy positiva y explicaron por qué. Quien observaba anotó lo que se veía, como dudas o intentos fallidos. Primero se probó el original y después el rediseño con las mismas personas. El protocolo completo está en datos/protocolo_prueba.md.")
doc.add_heading("3. Momentos de interacción", 1)
table(["Momento", "Fase", "Captura", "Qué hay en la pantalla"], [
    ["M1 Llegar y primera impresión", "Antes", "C02", "Banner rojo, cuenta regresiva 09:59 y colores muy fuertes"],
    ["M2 Explorar el mapa y ubicar la dirección", "Durante", "C03", "El pin rojo y el marcador 123 están en lugares distintos"],
    ["M3 Leer la dirección y el aviso", "Durante", "C04", "Texto rojo con presión y un error de ortografía"],
    ["M4 Decidir entre guardar y omitir", "Durante", "C05", "Guardar dirección en gris y OMITIR grande y verde"],
    ["M5 Terminar el paso y saber qué sigue", "Después", "C06", "No hay confirmación clara de lo que pasó"]], [5.2, 1.8, 1.6, 8.4])
IMG("evidencias/original/C01-pantalla-completa.png", 14.5, "Captura C01. Pantalla original con los elementos numerados.")
IMG("evidencias/original/C06-momento5-contexto-flujo.png", 16.5, "Captura C06. Contexto del flujo, de B-3 a B-5.")
doc.add_heading("4. Curva emocional del original", 1)
IMG("curva_emocional/curva_emocional.png", 16, f"Figura 1. Curva emocional con {npart} participantes y el promedio.")
P(f"El promedio general fue {pa.mean():+.1f}. El promedio más bajo fue el momento M{bajo}, {NOMBRES[bajo].lower()}, con {pa[bajo]:+.1f}. La mayor caída entre momentos seguidos llegó en M{caida}, {NOMBRES[caida].lower()}, con {delta:+.1f}. Esos son los dos momentos críticos.")
doc.add_heading("5. Matriz UX", 1)
ws = load_workbook("matriz/matriz_ux.xlsx")["Matriz UX"]; rows = []
for row in list(ws.iter_rows(values_only=True))[1:]:
    rows.append([row[0], row[2], row[4], row[5], row[6], f"{row[7]} {row[8]}", row[9]])
table(["Momento", "Emoción reportada", "Valoraciones", "Captura", "Evidencia observable", "Mejora y principio de IHC", "Validación"], rows, [2.4, 1.9, 2.0, 1.1, 3.0, 4.1, 2.6], 7.5)
P("Las valoraciones individuales van de -3 a +3. La columna de validación compara el promedio antes y después de cada momento. La matriz completa está en matriz/matriz_ux.xlsx.", size=9, italic=True)
doc.add_heading("6. Rediseño antes y después", 1)
P("Hicimos el rediseño con cinco cambios y lo armamos en Figma. La tarea y el contenido son los mismos que en el original.")
IMG("rediseño/antes_despues.png", 16.8, "Figura 2. Comparación visual antes y después.")
IMG("rediseño/figma_captura.png", 16.5, "Figura 3. Rediseño en Figma, pantalla Por confirmar.")
table(["N.º", "Cambio", "Emoción negativa que busca reducir", "Principio de IHC"], [
    ["1", "Sin banner rojo ni cuenta regresiva", "Presión y desconfianza en M1 y M3", "Heurística 8 de Nielsen. Diseño emocional de Norman"],
    ["2", "Un botón principal claro y Omitir como enlace", "Frustración en M4", "Heurística 3 y 5 de Nielsen. Jerarquía visual"],
    ["3", "Un solo pin con la dirección escrita", "Confusión en M2", "Heurística 1 y 4 de Nielsen"],
    ["4", "Colores calmados, buen contraste y textos corregidos", "Molestia en M1 y M3", "WCAG 1.4.3 de contraste"],
    ["5", "Barra de progreso con Paso 4 de 5 y confirmación", "Duda en M5", "Heurística 1 de Nielsen. Retroalimentación"]], [1.0, 5.2, 5.0, 5.8], 8.5)
doc.add_page_break()
doc.add_heading("7. Validación del rediseño", 1)
IMG("curva_emocional/curva_antes_despues.png", 16, "Figura 4. Promedio por momento antes y después.")
res = pd.read_csv("datos/resumen_comparacion.csv")
table(["Momento", "Promedio antes", "Promedio después", "Diferencia"], [[r["nombre"], f"{r['promedio_antes']:+.2f}", f"{r['promedio_despues']:+.2f}", f"{r['diferencia']:+.2f}"] for _, r in res.iterrows()], [7.2, 3.2, 3.4, 3.0], 9)
mejoran = int((pdv - pa > 0).sum())
P(f"El promedio mejoró en {mejoran} de los 5 momentos. Con {npart} personas no calculamos significancia, mostramos las diferencias tal como salieron.")
doc.add_heading("8. Preguntas de reflexión", 1)
P("¿Qué momento generó la emoción más negativa y qué evidencia lo demuestra?", True, after=2)
ev = list(dict.fromkeys(textos(a, bajo, "evidencia_observada") + [f"dijo {c}" for c in textos(a, bajo, "comentario")]))
cap = next(m["captura"] for m in MOMENTOS if m["id"] == bajo)
P(f"Fue M{bajo}, {NOMBRES[bajo].lower()}, con {pa[bajo]:+.1f} de promedio. Valores individuales, {', '.join(f'{p} {int(v):+d}' for p, v in va[bajo].items())}. "
  + ("Lo que se observó y dijeron, " + ". ".join(ev) + ". " if ev else "") + f"La captura de ese momento es {cap}.")
P("¿Cómo se relacionan las mejoras con la usabilidad, la accesibilidad y la confianza?", True, after=2)
P("Un botón principal claro mejora la usabilidad porque se entiende qué hacer. El contraste y los textos corregidos ayudan a la accesibilidad. Quitar la cuenta regresiva y la presión busca devolver confianza, porque el usuario decide sin sentirse apurado.")
P("¿Qué diferencia hubo entre la experiencia antes y después?", True, after=2)
mj = (pdv - pa).idxmax()
P(f"El promedio general pasó de {pa.mean():+.1f} a {pdv.mean():+.1f}. La mayor diferencia fue en M{mj}, {NOMBRES[mj].lower()}, con {(pdv - pa).max():+.1f} puntos.")
doc.add_heading("9. Conclusiones y limitaciones", 1)
P(f"En el original el momento más bajo fue M{bajo} y la mayor caída llegó en M{caida}. El rediseño cambió el promedio general de {pa.mean():+.1f} a {pdv.mean():+.1f}. Una mejora visual no basta, lo que la respalda son los números de las personas que probaron.")
P(f"Limitaciones. Son solo {npart} participantes. Las mismas personas ya habían visto el original cuando probaron el rediseño. Los resultados no se pueden generalizar.")
doc.add_heading("10. Enlaces y evidencias", 1)
P(f"Figma, {enl.get('figma', '')}")
P(f"Repositorio de GitHub, {enl.get('repositorio', '')}")
P("Evidencias en el repositorio. Capturas del original en evidencias/original. Hojas de prueba en evidencias/pruebas_usuarios. Matriz en matriz/matriz_ux.xlsx. Curva en curva_emocional. Contribuciones de cada integrante en TAREAS.md y en el historial de commits.")
os.makedirs("informe", exist_ok=True); doc.save(SALIDA); print("Guardado", SALIDA)
